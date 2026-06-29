"""Evaluate Hindi–Marathi language mix for question/answer pairs using a dedicated LLM judge."""

from __future__ import annotations

import argparse
import asyncio
import json
from pathlib import Path

import pandas as pd
from tqdm.asyncio import tqdm

from language_mix_evaluator import format_qa_pair, language_mix_agent

DEFAULT_MAX_CONCURRENT = 10
REPO_ROOT = Path(__file__).resolve().parent.parent


def resolve_path(path: Path) -> Path:
    if path.is_absolute():
        return path
    from_cwd = (Path.cwd() / path).resolve()
    if from_cwd.exists():
        return from_cwd
    return (REPO_ROOT / path).resolve()


def parse_args() -> argparse.Namespace:
    default_input = REPO_ROOT / "data" / "language_mix_eval" / "questions.json"
    default_output = REPO_ROOT / "data" / "language_mix_eval" / "language_quality_report.csv"
    default_json = REPO_ROOT / "data" / "language_mix_eval" / "evaluation.json"
    default_excel = REPO_ROOT / "data" / "language_mix_eval" / "language_quality_report.xlsx"

    parser = argparse.ArgumentParser(
        description="Evaluate Hindi–Marathi language mix for Q&A pairs."
    )
    parser.add_argument("--input", type=Path, default=default_input, help="Input CSV or JSON file")
    parser.add_argument("--output", type=Path, default=default_output, help="Output CSV report")
    parser.add_argument(
        "--eval-json",
        type=Path,
        default=default_json,
        help="Optional full evaluation JSON output",
    )
    parser.add_argument("--max-concurrent", type=int, default=DEFAULT_MAX_CONCURRENT)
    parser.add_argument(
        "--export-json",
        type=Path,
        default=None,
        metavar="JSON",
        help="Export an existing evaluation.json to CSV (no API calls)",
    )
    parser.add_argument(
        "--output-excel",
        type=Path,
        default=default_excel,
        help="Excel file to write/append results as a new sheet",
    )
    parser.add_argument(
        "--sheet-name",
        type=str,
        default="gemma evals",
        help="Sheet name for the new evaluation results in the Excel file",
    )
    return parser.parse_args()


def load_pairs(input_path: Path) -> list[dict[str, str]]:
    if input_path.suffix.lower() == ".json":
        with open(input_path, encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            data = data.get("items", data.get("data", []))
        return [{"question": item["question"], "answer": item["answer"]} for item in data]

    df = pd.read_csv(input_path)
    if "question" not in df.columns or "answer" not in df.columns:
        raise ValueError("Input CSV must contain 'question' and 'answer' columns.")
    return [
        {"question": str(row["question"]), "answer": str(row["answer"])}
        for _, row in df.iterrows()
        if pd.notna(row["question"]) and pd.notna(row["answer"])
    ]


def load_evaluation_json(json_path: Path) -> list[dict]:
    with open(json_path, encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, dict):
        data = data.get("items", data.get("data", []))
    return data


def export_evaluation_json_to_csv(json_path: Path, output_path: Path) -> int:
    items = load_evaluation_json(json_path)
    report_rows = [extract_report_row(item) for item in items]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(report_rows).to_csv(output_path, index=False, encoding="utf-8-sig")
    return len(report_rows)


def extract_report_row(item: dict) -> dict:
    evaluation = item.get("evaluation", {})
    lang = evaluation.get("dimensions", {}).get("language_quality", {})
    scores = lang.get("scores", {})
    mixed = evaluation.get("mixed_hindi_phrases", [])

    return {
        "question": item.get("question", ""),
        "answer": item.get("answer", ""),
        "summary": evaluation.get("summary", ""),
        "language_quality_avg": lang.get("average"),
        "mix_detected": evaluation.get("mix_detected"),
        "mix_severity": evaluation.get("mix_severity"),
        "mixed_hindi_phrases": "; ".join(mixed) if mixed else "",
        "grammar_score": scores.get("grammar", {}).get("score"),
        "grammar_evidence": scores.get("grammar", {}).get("evidence"),
        "marathi_terminology_score": scores.get("marathi_terminology", {}).get("score"),
        "marathi_terminology_evidence": scores.get("marathi_terminology", {}).get("evidence"),
        "language_purity_score": scores.get("language_purity", {}).get("score"),
        "language_purity_evidence": scores.get("language_purity", {}).get("evidence"),
        "fluency_score": scores.get("fluency", {}).get("score"),
        "fluency_evidence": scores.get("fluency", {}).get("evidence"),
    }


async def run_evaluation(
    pairs: list[dict[str, str]],
    max_concurrent: int,
) -> tuple[list[dict], list[dict]]:
    semaphore = asyncio.Semaphore(max_concurrent)

    async def evaluate_one(pair: dict[str, str]) -> dict | None:
        async with semaphore:
            item = {"question": pair["question"], "answer": pair["answer"]}
            message = format_qa_pair(pair["question"], pair["answer"])
            try:
                result = await language_mix_agent.run(message)
                item["evaluation"] = result.output.to_eval_dict()
                return item
            except Exception as exc:
                print(f"Failed: {pair['question'][:60]}... -> {exc}")
                return None

    tasks = [evaluate_one(pair) for pair in pairs]
    completed = await tqdm.gather(*tasks, desc="Evaluating language mix")

    full_results = [item for item in completed if item is not None]
    report_rows = [extract_report_row(item) for item in full_results]
    return full_results, report_rows


def write_excel_sheet(excel_path: Path, new_df: pd.DataFrame, sheet_name: str, existing_csv: Path | None) -> None:
    """Write new_df as sheet_name into excel_path, preserving any existing sheets."""
    from openpyxl import load_workbook

    excel_path.parent.mkdir(parents=True, exist_ok=True)

    if excel_path.exists():
        book = load_workbook(excel_path)
        # Remove sheet if it already exists so we can overwrite it cleanly
        if sheet_name in book.sheetnames:
            del book[sheet_name]
        with pd.ExcelWriter(excel_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            writer.book = book
            new_df.to_excel(writer, sheet_name=sheet_name, index=False)
    else:
        # New file — seed with existing CSV as first sheet if available
        with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
            if existing_csv and existing_csv.exists():
                pd.read_csv(existing_csv).to_excel(writer, sheet_name="previous evals", index=False)
            new_df.to_excel(writer, sheet_name=sheet_name, index=False)


def main() -> None:
    args = parse_args()
    args.input = resolve_path(args.input)
    args.output = resolve_path(args.output)
    if args.eval_json:
        args.eval_json = resolve_path(args.eval_json)
    if args.export_json:
        args.export_json = resolve_path(args.export_json)
    if args.output_excel:
        args.output_excel = resolve_path(args.output_excel)

    if args.export_json:
        count = export_evaluation_json_to_csv(args.export_json, args.output)
        print(f"Exported {count} rows from {args.export_json} to {args.output}")
        return

    pairs = load_pairs(args.input)
    if not pairs:
        raise SystemExit(f"No question/answer pairs found in {args.input}")

    print(f"Loaded {len(pairs)} pairs from {args.input}")
    full_results, report_rows = asyncio.run(
        run_evaluation(pairs, max_concurrent=args.max_concurrent)
    )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(report_rows).to_csv(args.output, index=False, encoding="utf-8-sig")
    print(f"Saved {len(report_rows)} rows to {args.output}")

    if args.output_excel:
        existing_csv = REPO_ROOT / "data" / "language_mix_eval" / "language_quality_report.csv"
        write_excel_sheet(args.output_excel, pd.DataFrame(report_rows), args.sheet_name, existing_csv)
        print(f"Saved sheet '{args.sheet_name}' to {args.output_excel}")

    if args.eval_json:
        args.eval_json.parent.mkdir(parents=True, exist_ok=True)
        with open(args.eval_json, "w", encoding="utf-8") as f:
            json.dump(full_results, f, ensure_ascii=False, indent=2)
        print(f"Saved full evaluation JSON to {args.eval_json}")


if __name__ == "__main__":
    main()
